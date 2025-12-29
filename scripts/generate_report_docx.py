from __future__ import annotations

from pathlib import Path

from docx import Document


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out_dir = root / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "final_report.docx"

    doc = Document()
    doc.add_heading("MLOps Assignment-1: Heart Disease UCI — End-to-End MLOps", level=0)

    doc.add_paragraph(
        "This document is generated from the repository implementation. Replace placeholders "
        "with screenshots and final deployment proof as required by the assignment."
    )

    sections = [
        "1. Setup / Install Instructions",
        "2. Data Acquisition, Cleaning, and EDA",
        "3. Feature Engineering and Model Development",
        "4. Experiment Tracking (MLflow)",
        "5. Model Packaging and Reproducibility",
        "6. CI/CD Pipeline and Automated Testing",
        "7. Containerization (Docker)",
        "8. Production Deployment (Kubernetes)",
        "9. Monitoring and Logging",
        "10. Repository Link and Artifacts",
    ]

    for i, title in enumerate(sections):
        doc.add_heading(title, level=1)
        doc.add_paragraph(
            "Fill in the following for this section:\n"
            "- Summary (what was built)\n"
            "- Commands used\n"
            "- Key results/metrics\n"
            "- Screenshots (paste from screenshots/ folder)\n"
        )
        if i < len(sections) - 1:
            doc.add_page_break()

    doc.add_page_break()
    doc.add_heading("Appendix: Commands", level=1)
    doc.add_paragraph("Pipeline: python -m heart_disease_mlops.pipeline run")
    doc.add_paragraph("API: uvicorn heart_disease_mlops.serving.app:app --host 0.0.0.0 --port 8000")
    doc.add_paragraph("Docker: docker build -t heart-mlops:latest .")
    doc.add_paragraph("K8s: kubectl apply -f deploy/k8s/api.yaml")
    doc.add_paragraph("Monitoring: cd deploy/monitoring ; docker compose up --build")

    doc.save(out_path)
    print(f"Wrote: {out_path}")


if __name__ == "__main__":
    main()
